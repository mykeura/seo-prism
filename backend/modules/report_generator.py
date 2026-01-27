"""
SEO Report Generator Module

Generates professional PowerPoint presentations with SEO analysis results.
Supports bilingual reports (Spanish/English) with corporate design.
"""

from typing import Dict, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from datetime import datetime
from .report_translations import get_translation, get_impact_template


class SEOReportGenerator:
    """Generates professional SEO reports in PowerPoint format."""
    
    # Corporate color theme (easily customizable)
    THEME_COLORS = {
        'primary': RGBColor(30, 58, 138),      # Dark blue
        'secondary': RGBColor(100, 116, 139),  # Slate gray
        'accent': RGBColor(59, 130, 246),      # Bright blue
        'success': RGBColor(16, 185, 129),     # Green (Grade A)
        'warning': RGBColor(245, 158, 11),     # Orange (Grade C-D)
        'danger': RGBColor(239, 68, 68),       # Red (Grade F)
        'background': RGBColor(255, 255, 255), # White
        'text': RGBColor(31, 41, 55),          # Dark gray
        'text_light': RGBColor(107, 114, 128), # Light gray
        'header_bg': RGBColor(30, 58, 138),    # Header background
    }
    
    def __init__(self, custom_theme: Dict = None):
        """
        Initialize the SEO report generator.
        
        Args:
            custom_theme: Optional dictionary of custom colors to override defaults
        """
        self.theme = self.THEME_COLORS.copy()
        if custom_theme:
            self.theme.update(custom_theme)
    
    def generate_report(self, results: Dict, output_path: str, lang: str = 'en') -> str:
        """
        Generate a professional SEO report in PowerPoint format.
        
        Args:
            results: Scan results dictionary with scan info, issues, pages, seo_grade, resource_analysis
            output_path: Path where the PowerPoint file will be saved
            lang: Language code ('es' or 'en')
        
        Returns:
            Path to the generated PowerPoint file
        """
        # Create presentation
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)
        
        # Generate slides
        self._add_cover_slide(prs, results, lang)
        self._add_executive_summary(prs, results, lang)
        self._add_impact_slide(prs, results, lang)
        self._add_distribution_slide(prs, results, lang)
        self._add_detailed_analysis_slides(prs, results, lang)
        self._add_recommendations_slide(prs, results, lang)
        self._add_notes_slide(prs, lang)
        
        # Save presentation
        prs.save(output_path)
        return output_path
    
    def _add_cover_slide(self, prs: Presentation, results: Dict, lang: str):
        """Add cover slide with logo placeholder, grade, and key info."""
        slide_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(slide_layout)
        
        # Add gradient background
        background = slide.background
        fill = background.fill
        fill.gradient()
        fill.gradient_angle = 90
        fill.gradient_stops[0].color.rgb = self.theme['primary']
        fill.gradient_stops[1].color.rgb = self.theme['accent']
        
        # Add logo placeholder
        logo_shape = slide.shapes.add_shape(
            1,  # Rectangle
            Inches(1), Inches(0.5), Inches(2), Inches(1.5)
        )
        logo_shape.fill.solid()
        logo_shape.fill.fore_color.rgb = self.theme['background']
        logo_shape.line.color.rgb = self.theme['text_light']
        
        logo_text = logo_shape.text_frame
        logo_text.text = get_translation(lang, 'placeholder_logo')
        logo_text.paragraphs[0].alignment = PP_ALIGN.CENTER
        for paragraph in logo_text.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(10)
                run.font.color.rgb = self.theme['text_light']
        
        # Add title
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(2.5), Inches(9), Inches(1)
        )
        title_frame = title_box.text_frame
        title_frame.text = get_translation(lang, 'report_title')
        title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        title_frame.paragraphs[0].font.size = Pt(44)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        # Add client name placeholder
        client_box = slide.shapes.add_textbox(
            Inches(2), Inches(3.5), Inches(6), Inches(0.5)
        )
        client_frame = client_box.text_frame
        client_frame.text = get_translation(lang, 'placeholder_client_name')
        client_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        client_frame.paragraphs[0].font.size = Pt(18)
        client_frame.paragraphs[0].font.color.rgb = RGBColor(200, 200, 200)
        
        # Add URL
        url_box = slide.shapes.add_textbox(
            Inches(2), Inches(4.1), Inches(6), Inches(0.4)
        )
        url_frame = url_box.text_frame
        url_frame.text = f"{get_translation(lang, 'analyzed_url')}: {results.get('scan', {}).get('url', 'N/A')}"
        url_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        url_frame.paragraphs[0].font.size = Pt(14)
        url_frame.paragraphs[0].font.color.rgb = RGBColor(180, 180, 180)
        
        # Add date
        date_box = slide.shapes.add_textbox(
            Inches(2), Inches(4.6), Inches(6), Inches(0.4)
        )
        date_frame = date_box.text_frame
        date_frame.text = f"{get_translation(lang, 'analysis_date')}: {results.get('scan', {}).get('timestamp', datetime.now().strftime('%Y-%m-%d'))}"
        date_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        date_frame.paragraphs[0].font.size = Pt(14)
        date_frame.paragraphs[0].font.color.rgb = RGBColor(180, 180, 180)
        
        # Add SEO Grade box (large and prominent)
        grade_data = results.get('seo_grade', {})
        grade = grade_data.get('grade', 'N/A')
        score = grade_data.get('score', 0)
        
        # Determine grade color
        grade_color = self.theme['success']
        if grade == 'B':
            grade_color = self.theme['accent']
        elif grade == 'C':
            grade_color = self.theme['warning']
        elif grade == 'D':
            grade_color = RGBColor(255, 140, 0)  # Dark orange
        elif grade == 'F':
            grade_color = self.theme['danger']
        
        grade_box = slide.shapes.add_shape(
            1,  # Rectangle
            Inches(3.5), Inches(5.2), Inches(3), Inches(2)
        )
        grade_box.fill.solid()
        grade_box.fill.fore_color.rgb = grade_color
        grade_box.line.color.rgb = self.theme['background']
        grade_box.line.width = Pt(3)
        
        grade_text = grade_box.text_frame
        grade_text.word_wrap = True
        
        # Add "GRADE" label
        p_grade_label = grade_text.paragraphs[0]
        p_grade_label.text = get_translation(lang, 'seo_grade').upper()
        p_grade_label.alignment = PP_ALIGN.CENTER
        p_grade_label.font.size = Pt(12)
        p_grade_label.font.bold = True
        p_grade_label.font.color.rgb = RGBColor(255, 255, 255)
        p_grade_label.space_after = Pt(5)
        
        # Add grade letter (large)
        p_grade = grade_text.add_paragraph()
        p_grade.text = grade
        p_grade.alignment = PP_ALIGN.CENTER
        p_grade.font.size = Pt(72)
        p_grade.font.bold = True
        p_grade.font.color.rgb = RGBColor(255, 255, 255)
        p_grade.space_after = Pt(5)
        
        # Add score
        p_score = grade_text.add_paragraph()
        p_score.text = f"{get_translation(lang, 'score').upper()}: {score}%"
        p_score.alignment = PP_ALIGN.CENTER
        p_score.font.size = Pt(16)
        p_score.font.bold = True
        p_score.font.color.rgb = RGBColor(255, 255, 255)
    
    def _add_executive_summary(self, prs: Presentation, results: Dict, lang: str):
        """Add executive summary slide with key metrics."""
        slide_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(slide_layout)
        
        # Add header bar
        header_shape = slide.shapes.add_shape(
            1,  # Rectangle
            Inches(0), Inches(0), Inches(10), Inches(1)
        )
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = self.theme['primary']
        header_shape.line.fill.background()
        
        # Add title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
        title_frame = title_box.text_frame
        title_frame.text = get_translation(lang, 'executive_summary')
        title_frame.paragraphs[0].font.size = Pt(32)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        # Add grade explanation
        grade_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.2), Inches(9), Inches(0.6)
        )
        grade_frame = grade_box.text_frame
        grade_frame.text = get_translation(lang, 'grade_explanation')
        grade_frame.paragraphs[0].font.size = Pt(12)
        grade_frame.paragraphs[0].font.color.rgb = self.theme['text_light']
        
        # Get metrics
        scan = results.get('scan', {})
        grade_data = results.get('seo_grade', {})
        breakdown = grade_data.get('breakdown', {})
        
        # Add metric cards
        metrics = [
            {
                'label': get_translation(lang, 'total_pages'),
                'value': str(scan.get('total_pages', 0)),
                'color': self.theme['accent']
            },
            {
                'label': get_translation(lang, 'total_issues'),
                'value': str(scan.get('total_issues', 0)),
                'color': self.theme['warning']
            },
            {
                'label': get_translation(lang, 'high_priority'),
                'value': str(breakdown.get('high', 0)),
                'color': self.theme['danger']
            },
            {
                'label': get_translation(lang, 'medium_priority'),
                'value': str(breakdown.get('medium', 0)),
                'color': self.theme['warning']
            },
        ]
        
        # Position metric cards
        card_width = Inches(2.2)
        card_height = Inches(1.2)
        card_spacing = Inches(0.2)
        start_x = Inches(0.5)
        start_y = Inches(2.0)
        
        for i, metric in enumerate(metrics):
            x = start_x + i * (card_width + card_spacing)
            
            # Add card background
            card = slide.shapes.add_shape(1, x, start_y, card_width, card_height)
            card.fill.solid()
            card.fill.fore_color.rgb = self.theme['background']
            card.line.color.rgb = metric['color']
            card.line.width = Pt(2)
            
            # Add value
            value_box = slide.shapes.add_textbox(x, start_y + Inches(0.2), card_width, Inches(0.5))
            value_frame = value_box.text_frame
            value_frame.text = metric['value']
            value_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            value_frame.paragraphs[0].font.size = Pt(28)
            value_frame.paragraphs[0].font.bold = True
            value_frame.paragraphs[0].font.color.rgb = metric['color']
            
            # Add label
            label_box = slide.shapes.add_textbox(x, start_y + Inches(0.7), card_width, Inches(0.4))
            label_frame = label_box.text_frame
            label_frame.text = metric['label']
            label_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            label_frame.paragraphs[0].font.size = Pt(10)
            label_frame.paragraphs[0].font.color.rgb = self.theme['text_light']
        
        # Add critical issues summary
        issues = results.get('issues', [])
        critical_issues = [i for i in issues if i.get('severity') == 'high'][:5]
        
        if critical_issues:
            critical_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(3.5), Inches(9), Inches(3.5)
            )
            critical_frame = critical_box.text_frame
            critical_frame.text = get_translation(lang, 'critical_issues')
            critical_frame.paragraphs[0].font.size = Pt(16)
            critical_frame.paragraphs[0].font.bold = True
            critical_frame.paragraphs[0].font.color.rgb = self.theme['danger']
            critical_frame.paragraphs[0].space_after = Pt(10)
            
            for issue in critical_issues:
                p = critical_frame.add_paragraph()
                p.text = f"• {issue.get('issue_type', 'Unknown')}: {issue.get('url', 'N/A')}"
                p.font.size = Pt(11)
                p.font.color.rgb = self.theme['text']
                p.space_after = Pt(5)
    
    def _add_impact_slide(self, prs: Presentation, results: Dict, lang: str):
        """Add slide explaining why these issues matter for Google."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # Add header bar
        header_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = self.theme['primary']
        header_shape.line.fill.background()
        
        # Add title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
        title_frame = title_box.text_frame
        title_frame.text = get_translation(lang, 'impact_on_google')
        title_frame.paragraphs[0].font.size = Pt(32)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        # Add content
        content_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.2), Inches(9), Inches(6)
        )
        content_frame = content_box.text_frame
        
        # Introduction
        p_intro = content_frame.paragraphs[0]
        p_intro.text = get_translation(lang, 'issues_by_severity')
        p_intro.font.size = Pt(14)
        p_intro.font.bold = True
        p_intro.font.color.rgb = self.theme['primary']
        p_intro.space_after = Pt(15)
        
        # Severity explanations
        severity_info = [
            (get_translation(lang, 'high_priority'), self.theme['danger'], get_translation(lang, 'must_correct')),
            (get_translation(lang, 'medium_priority'), self.theme['warning'], get_translation(lang, 'should_correct')),
            (get_translation(lang, 'low_priority'), self.theme['success'], get_translation(lang, 'could_correct')),
        ]
        
        for severity, color, action in severity_info:
            p = content_frame.add_paragraph()
            p.text = f"{severity}: {action}"
            p.font.size = Pt(12)
            p.font.color.rgb = color
            p.font.bold = True
            p.space_before = Pt(10)
            p.space_after = Pt(8)
    
    def _add_distribution_slide(self, prs: Presentation, results: Dict, lang: str):
        """Add slide with issue distribution by category."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # Add header bar
        header_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = self.theme['primary']
        header_shape.line.fill.background()
        
        # Add title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
        title_frame = title_box.text_frame
        title_frame.text = get_translation(lang, 'distribution')
        title_frame.paragraphs[0].font.size = Pt(32)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        # Group issues by category
        issues = results.get('issues', [])
        categories = {}
        for issue in issues:
            issue_type = issue.get('issue_type', 'unknown')
            if issue_type not in categories:
                categories[issue_type] = []
            categories[issue_type].append(issue)
        
        # Sort by count (descending)
        sorted_categories = sorted(categories.items(), key=lambda x: len(x[1]), reverse=True)
        
        # Create table
        table = slide.shapes.add_table(
            rows=min(len(sorted_categories) + 1, 10),  # Max 10 rows
            cols=2,
            left=Inches(0.5),
            top=Inches(1.2),
            width=Inches(9),
            height=Inches(6)
        ).table
        
        # Style table
        table.style = 'Medium Style 2 - Accent 1'
        
        # Header row
        table.cell(0, 0).text = get_translation(lang, 'description')
        table.cell(0, 1).text = 'Count'
        
        # Fill table
        for idx, (category, cat_issues) in enumerate(sorted_categories[:9]):  # Max 9 data rows
            if idx + 1 >= len(table.rows):
                break
            table.cell(idx + 1, 0).text = category
            table.cell(idx + 1, 1).text = str(len(cat_issues))
        
        # Style cells
        for row in table.rows:
            for cell in row.cells:
                cell.text_frame.paragraphs[0].font.size = Pt(11)
                cell.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    
    def _add_detailed_analysis_slides(self, prs: Presentation, results: Dict, lang: str):
        """Add detailed analysis slides for each issue category."""
        issues = results.get('issues', [])
        
        # Group issues by category
        categories = {}
        for issue in issues:
            issue_type = issue.get('issue_type', 'unknown')
            if issue_type not in categories:
                categories[issue_type] = []
            categories[issue_type].append(issue)
        
        # Only add slides for categories with issues
        for category, cat_issues in sorted(categories.items()):
            if len(cat_issues) == 0:
                continue
            
            self._add_category_slide(prs, category, cat_issues, lang)
    
    def _add_category_slide(self, prs: Presentation, category: str, issues: List[Dict], lang: str):
        """Add a slide for a specific issue category."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # Add header bar
        header_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = self.theme['primary']
        header_shape.line.fill.background()
        
        # Add title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
        title_frame = title_box.text_frame
        title_frame.text = f"{category.upper()} ({len(issues)} issues)"
        title_frame.paragraphs[0].font.size = Pt(28)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        # Get impact template
        impact_template = get_impact_template(lang, category)
        
        # Add impact
        impact_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.2), Inches(9), Inches(1.2)
        )
        impact_frame = impact_box.text_frame
        impact_frame.text = f"{get_translation(lang, 'impact_on_google')}:\n{impact_template.get('impact', '')}"
        impact_frame.paragraphs[0].font.size = Pt(11)
        impact_frame.paragraphs[0].font.bold = True
        impact_frame.paragraphs[0].font.color.rgb = self.theme['text']
        impact_frame.paragraphs[1].font.size = Pt(10)
        impact_frame.paragraphs[1].font.color.rgb = self.theme['text_light']
        
        # Create table for issues
        num_rows = min(len(issues) + 1, 8)  # Max 8 rows total
        table = slide.shapes.add_table(
            rows=num_rows,
            cols=3,
            left=Inches(0.5),
            top=Inches(2.6),
            width=Inches(9),
            height=Inches(4)
        ).table
        
        # Header row
        table.cell(0, 0).text = get_translation(lang, 'description')
        table.cell(0, 1).text = get_translation(lang, 'examples')
        table.cell(0, 2).text = get_translation(lang, 'priority')
        
        # Fill table
        for idx, issue in enumerate(issues[:7]):  # Max 7 data rows
            if idx + 1 >= len(table.rows):
                break
            
            # Description
            table.cell(idx + 1, 0).text = issue.get('description', 'N/A')
            
            # Examples (URL)
            table.cell(idx + 1, 1).text = issue.get('url', 'N/A')
            
            # Priority
            priority = issue.get('severity', 'medium').upper()
            table.cell(idx + 1, 2).text = priority
            
            # Color based on priority
            if priority == 'HIGH':
                table.cell(idx + 1, 2).fill.solid()
                table.cell(idx + 1, 2).fill.fore_color.rgb = self.theme['danger']
                table.cell(idx + 1, 2).text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
            elif priority == 'MEDIUM':
                table.cell(idx + 1, 2).fill.solid()
                table.cell(idx + 1, 2).fill.fore_color.rgb = self.theme['warning']
                table.cell(idx + 1, 2).text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
            else:
                table.cell(idx + 1, 2).fill.solid()
                table.cell(idx + 1, 2).fill.fore_color.rgb = self.theme['success']
                table.cell(idx + 1, 2).text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        
        # Style cells
        for row in table.rows:
            for cell in row.cells:
                cell.text_frame.paragraphs[0].font.size = Pt(9)
                cell.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
                cell.text_frame.word_wrap = True
        
        # Add priority note
        priority_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(6.7), Inches(9), Inches(0.7)
        )
        priority_frame = priority_box.text_frame
        priority_frame.text = impact_template.get('priority_note', '')
        priority_frame.paragraphs[0].font.size = Pt(9)
        priority_frame.paragraphs[0].font.color.rgb = self.theme['text_light']
    
    def _add_recommendations_slide(self, prs: Presentation, results: Dict, lang: str):
        """Add recommendations and next steps slide."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # Add header bar
        header_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = self.theme['primary']
        header_shape.line.fill.background()
        
        # Add title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
        title_frame = title_box.text_frame
        title_frame.text = get_translation(lang, 'next_steps')
        title_frame.paragraphs[0].font.size = Pt(32)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        # Add recommendations
        content_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.2), Inches(9), Inches(6)
        )
        content_frame = content_box.text_frame
        
        # Priority order
        recommendations = [
            get_translation(lang, 'must_correct'),
            get_translation(lang, 'should_correct'),
            get_translation(lang, 'could_correct'),
        ]
        
        for i, rec in enumerate(recommendations):
            p = content_frame.add_paragraph()
            p.text = f"{i + 1}. {rec}"
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = self.theme['text']
            p.space_before = Pt(8)
            p.space_after = Pt(8)
        
        # Add call to action
        cta_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(5.5), Inches(9), Inches(1.5)
        )
        cta_frame = cta_box.text_frame
        cta_frame.text = get_translation(lang, 'call_to_action')
        cta_frame.paragraphs[0].font.size = Pt(11)
        cta_frame.paragraphs[0].font.italic = True
        cta_frame.paragraphs[0].font.color.rgb = self.theme['accent']
    
    def _add_notes_slide(self, prs: Presentation, lang: str):
        """Add editable notes slide."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # Add header bar
        header_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = self.theme['primary']
        header_shape.line.fill.background()
        
        # Add title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
        title_frame = title_box.text_frame
        title_frame.text = get_translation(lang, 'notes')
        title_frame.paragraphs[0].font.size = Pt(32)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        # Add placeholder for notes
        notes_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.2), Inches(9), Inches(6)
        )
        notes_frame = notes_box.text_frame
        notes_frame.text = get_translation(lang, 'placeholder_notes')
        notes_frame.paragraphs[0].font.size = Pt(11)
        notes_frame.paragraphs[0].font.color.rgb = self.theme['text_light']
        notes_frame.paragraphs[0].font.italic = True
